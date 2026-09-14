from flask import Flask, render_template, request, jsonify, send_file
from modules.transcript_analyzer import analyze_transcript
from modules.proposal_retriever import find_similar_proposal
from modules.proposal_generator import generate_proposal
import os

app = Flask(__name__)


# --------------------------------------------------
# HOME PAGE
# --------------------------------------------------

@app.route("/")
def home():
    return render_template("index.html")


# --------------------------------------------------
# GENERATE PROPOSAL
# --------------------------------------------------

@app.route("/generate", methods=["POST"])
def generate():

    transcript = request.form.get("transcript", "").strip()

    if not transcript:
        return jsonify({
            "error": "Please enter a sales call transcript."
        }), 400

    try:

        # 1. Analyze the sales call
        analysis = analyze_transcript(transcript)

        # 2. Find the most similar previous proposal
        proposal_file, similarity_score = find_similar_proposal(
            analysis
        )

        # 3. Load the previous proposal
        proposal_path = os.path.join(
            "data",
            "past_proposals",
            proposal_file
        )

        with open(
            proposal_path,
            "r",
            encoding="utf-8"
        ) as file:
            past_proposal = file.read()

        # 4. Generate personalized proposal
        proposal = generate_proposal(
            analysis,
            past_proposal
        )

        # 5. Return result to frontend
        return jsonify({
            "analysis": analysis,
            "proposal": proposal,
            "proposal_file": proposal_file,
            "similarity_score": round(
                float(similarity_score),
                2
            )
        })

    except Exception as e:

        print("ERROR:", e)

        return jsonify({
            "error": str(e)
        }), 500


# --------------------------------------------------
# DOWNLOAD DOCX
# --------------------------------------------------

@app.route("/download-docx", methods=["POST"])
def download_docx():

    data = request.get_json()

    if not data:
        return jsonify({
            "error": "No data received."
        }), 400

    proposal = data.get("proposal", "").strip()

    if not proposal:
        return jsonify({
            "error": "No proposal available."
        }), 400

    from docx import Document

    document = Document()

    document.add_heading(
        "Project Proposal",
        level=0
    )

    for paragraph in proposal.split("\n"):

        paragraph = paragraph.strip()

        if paragraph:
            document.add_paragraph(paragraph)

    # Make sure output folder exists
    os.makedirs("output", exist_ok=True)

    output_path = os.path.join(
        "output",
        "personalized_proposal.docx"
    )

    document.save(output_path)

    return send_file(
        output_path,
        as_attachment=True,
        download_name="personalized_proposal.docx"
    )


# --------------------------------------------------
# DOWNLOAD PDF
# --------------------------------------------------

@app.route("/download-pdf", methods=["POST"])
def download_pdf():

    data = request.get_json()

    if not data:
        return jsonify({
            "error": "No data received."
        }), 400

    proposal = data.get("proposal", "").strip()

    if not proposal:
        return jsonify({
            "error": "No proposal available."
        }), 400

    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer
    )

    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm

    # Make sure output folder exists
    os.makedirs("output", exist_ok=True)

    output_path = os.path.join(
        "output",
        "personalized_proposal.pdf"
    )

    document = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm
    )

    styles = getSampleStyleSheet()

    story = []

    story.append(
        Paragraph(
            "Project Proposal",
            styles["Title"]
        )
    )

    story.append(
        Spacer(1, 12)
    )

    for paragraph in proposal.split("\n"):

        paragraph = paragraph.strip()

        if paragraph:

            # Escape characters that have special
            # meaning in ReportLab Paragraph
            paragraph = (
                paragraph
                .replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )

            story.append(
                Paragraph(
                    paragraph,
                    styles["BodyText"]
                )
            )

            story.append(
                Spacer(1, 6)
            )

    document.build(story)

    return send_file(
        output_path,
        as_attachment=True,
        download_name="personalized_proposal.pdf"
    )


# --------------------------------------------------
# START FLASK APPLICATION
# --------------------------------------------------

if __name__ == "__main__":
    app.run(
        debug=True
    )